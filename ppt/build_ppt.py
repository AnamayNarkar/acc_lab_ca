"""Build ppt/CostScale.pptx (16:9) from results/*.json and report/figures. Speaker notes on every slide."""
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
ROOT = Path(__file__).resolve().parent.parent
R = lambda n: json.load(open(ROOT / "results" / n))
DS, SIM, FC, CAL, NV, WL = R("dataset_summary.json"), R("sim_results.json"), R("forecast_results.json"), R("calibration.json"), R("newsvendor_check.json"), R("live_workload.json")
SIM16 = R("sim_results_assumed_cap16.json")
LIVE = R("live_results.json") if (ROOT / "results/live_results.json").exists() else {}
T = SIM["ec2"]["test"]; FIG = ROOT / "report/figures"
K = lambda x: f"${x/1000:,.1f}k"
BLUE, ORANGE, AQUA = RGBColor(0x2a, 0x78, 0xd6), RGBColor(0xeb, 0x68, 0x34), RGBColor(0x1b, 0xaf, 0x7a)
INK, MUTED, BG = RGBColor(0x0b, 0x0b, 0x0b), RGBColor(0x52, 0x51, 0x4e), RGBColor(0xfc, 0xfc, 0xfb)
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]

def slide(title, notes, kicker=None):
    s = prs.slides.add_slide(blank)
    bg = s.background.fill; bg.solid(); bg.fore_color.rgb = BG
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), prs.slide_height); bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background()
    tb = s.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.2), Inches(0.9)); tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; r = p.add_run(); r.text = title; r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = INK; r.font.name = "Calibri"
    if kicker:
        p2 = tf.add_paragraph(); r2 = p2.add_run(); r2.text = kicker; r2.font.size = Pt(15); r2.font.color.rgb = MUTED; r2.font.name = "Calibri"
    s.notes_slide.notes_text_frame.text = notes
    return s

def bullets(s, items, x=0.6, y=1.7, w=6.2, h=5.2, size=18, gap=8):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.space_after = Pt(gap)
        sub = it.startswith("  "); txt = it.strip()
        r = p.add_run(); r.text = ("– " if sub else "• ") + txt; r.font.size = Pt(size - 3 if sub else size); r.font.color.rgb = MUTED if sub else INK; r.font.name = "Calibri"
        if sub: p.level = 1
    return tb

def pic(s, name, x, y, w=None, h=None):
    return s.shapes.add_picture(str(FIG / name), Inches(x), Inches(y), width=Inches(w) if w else None, height=Inches(h) if h else None)

def stat(s, x, y, big, small, color=BLUE, w=2.9):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(1.4)); tf = tb.text_frame; tf.word_wrap = True
    r = tf.paragraphs[0].add_run(); r.text = big; r.font.size = Pt(36); r.font.bold = True; r.font.color.rgb = color; r.font.name = "Calibri"
    p = tf.add_paragraph(); r = p.add_run(); r.text = small; r.font.size = Pt(13); r.font.color.rgb = MUTED; r.font.name = "Calibri"

def table(s, rows, x, y, w, colw, size=12, h=0.4):
    sh = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h * len(rows))); t = sh.table
    for j, cw in enumerate(colw): t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.text = str(v); p = c.text_frame.paragraphs[0]; p.runs[0].font.size = Pt(size); p.runs[0].font.name = "Calibri"
            p.runs[0].font.bold = (i == 0); p.runs[0].font.color.rgb = INK
            c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0xd9, 0xd9, 0xd9) if i == 0 else RGBColor(0xff, 0xff, 0xff)
            c.margin_top = c.margin_bottom = Emu(30000)
    return t

# 1 title
s = slide("", "Introduce the project in one sentence: a resource manager that keeps a cheap base of EC2 sized by a forecast and sends the rest to Lambda. Mention it is a university cloud-computing project, built and measured on AWS in Mumbai, using a real production trace from Azure. Names and IDs are to be filled in.")
tt = s.shapes.add_textbox(Inches(0.9), Inches(1.4), Inches(11.5), Inches(1.6)); rr = tt.text_frame.paragraphs[0].add_run(); rr.text = "CostScale"; rr.font.size = Pt(66); rr.font.bold = True; rr.font.color.rgb = INK; rr.font.name = "Calibri"
tt2 = s.shapes.add_textbox(Inches(0.9), Inches(2.9), Inches(11.5), Inches(0.9)); tt2.text_frame.word_wrap = True; r2 = tt2.text_frame.paragraphs[0].add_run(); r2.text = "ML-driven hybrid container and serverless resource management on AWS"; r2.font.size = Pt(26); r2.font.color.rgb = BLUE; r2.font.name = "Calibri"
bullets(s, ["[TO FILL: student names and IDs]", "[TO FILL: programme, institute, course]", "October 2026"], x=0.9, y=4.3, size=20)
# 2 problem
s = slide("The problem", "Two ways to buy compute: reserved containers or VMs are cheap per unit but need lead time; Lambda is instant and pay per use but costs several times more per unit of work. Traffic varies, so sizing for peak wastes money and reacting late drops or delays requests. The task: predict workload and allocate resources. Rubric items: inputs and features, algorithm and horizon, allocation mechanism, evaluation on accuracy, utilisation, latency, cost and energy, and what happens when the prediction is wrong.", "Predict workload, then allocate resources")
bullets(s, ["Containers / VMs: cheap per unit, but slow to provision", "Lambda: instant, pay per use, several times costlier per unit", "Load varies; most functions are idle, a few are very busy", "Goal: lowest cost while serving every request in time"], w=6.4)
stat(s, 7.6, 1.9, f"{DS['heavy_functions']}", "heavy Azure functions simulated\n(>= $50/month on training days)", BLUE)
stat(s, 7.6, 3.6, f"{DS['functions_on_all_14_days']:,}", "functions present on all 14 days\n(of 72k in the trace)", ORANGE)
# 3 why autoscaling not enough
r7, r3 = T["3_reactive_u0.7"], T["3_reactive_u0.3"]
s = slide("Why plain autoscaling is not enough", f"Reactive autoscaling sizes capacity from load it saw five minutes ago at a target utilisation u. In the simulation, every u we tried leaves load unserved because there is no fallback: u 0.7 costs {K(r7['cost'])} and drops {r7['unserved_pct']:.1f} percent; u 0.3 costs {K(r3['cost'])} and still drops {r3['unserved_pct']:.1f} percent. The hybrid in the bottom-left corner of the chart serves everything for {K(T['6_hybrid_rolling']['cost'])}. The seasonal line is an emulation of predictive-scaling style capacity from the same hour on previous days, not the AWS algorithm.", "Every reactive setting drops load or overpays")
pic(s, "unserved_vs_cost.png", 0.4, 1.7, w=8.0)
bullets(s, [f"Reactive u=0.7: {K(r7['cost'])}, {r7['unserved_pct']:.1f}% load dropped", f"Reactive u=0.3: {K(r3['cost'])}, still {r3['unserved_pct']:.1f}% dropped", f"Hybrid-rolling: {K(T['6_hybrid_rolling']['cost'])}, 0% dropped", "Test days, 101 heavy functions, AWS Mumbai prices"], x=8.6, y=2.0, w=4.5, size=17)
# 4 idea
s = slide("The idea: a base plus overflow", f"Keep a base of always-on EC2 capacity. Everything above the base goes to Lambda. The base is the quantile of forecast load whose level is q-star equals one minus f over L, where f is the container cost and L the Lambda cost of the same unit of load for a minute. If L is not above f the base is zero. We checked the rule against brute-force search per function and hour: the difference is {NV['relative_gap_pct']:.4f} percent. Forecasting sets the slow thing, instance count; a per-minute loop sets ALB weights.", "Newsvendor: base = q*-quantile of forecast load")
bullets(s, ["Base: EC2 Auto Scaling Group (cheap, always on)", "Overflow: AWS Lambda (instant, pay per use)", "ALB weighted target groups split the traffic", "Forecast sets instance count (hourly); observed rate sets weights (per minute)"], w=6.2)
tb = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.2), Inches(2.0), Inches(5.6), Inches(2.6)); tb.fill.solid(); tb.fill.fore_color.rgb = RGBColor(0xe8, 0xf0, 0xfb); tb.line.color.rgb = BLUE
tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
for i, (txt, sz, b) in enumerate([("base  b* = F⁻¹(q*)", 26, True), ("q* = 1 − f / L", 26, True), ("f: container $ per unit-minute\nL: Lambda $ per unit-minute", 15, False)]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = txt; r.font.size = Pt(sz); r.font.bold = b; r.font.color.rgb = INK; r.font.name = "Calibri"
stat(s, 7.2, 5.0, f"{abs(NV['relative_gap_pct']):.4f}%", "gap between the rule and brute-force search (test hours)", AQUA, w=5.6)
# 5 data
s = slide("Data", f"Azure Functions 2019 trace, CC-BY, from Shahrad et al., USENIX ATC 2020. Fourteen days of per-minute invocations, with durations and memory. Most functions are rarely invoked, the median function on day one gets {DS['median_invocations_per_day_day1']:.0f} invocations, so we filter. We keep functions present on all 14 days with at least 5 requests per minute: {DS['functions_avg_ge_5rpm_all14days']:,}, of which {DS['functions_with_traffic_duration_memory']:,} also have duration and memory. A function is heavy if its cheaper single-platform option costs at least 50 dollars per month on training days. Load is concurrency: invocations per minute over 60 times average duration. Azure traffic is priced with AWS prices. Split is chronological per function: days 1-10 train, 11-12 validation, 13-14 test.", "Azure Functions 2019 trace, filtered to busy functions")
table(s, [["Step", "Count"], ["Distinct functions (14 days)", f"{DS['functions_total_distinct']:,}"], ["On day 1", f"{DS['day1_functions']:,}"], ["On all 14 days", f"{DS['functions_on_all_14_days']:,}"], [">= 5 req/min, all 14 days", f"{DS['functions_avg_ge_5rpm_all14days']:,}"], ["+ duration + memory", f"{DS['functions_with_traffic_duration_memory']:,}"], ["Heavy (>= $50/mo)", f"{DS['heavy_functions']}"]], 0.6, 1.8, 6.2, [4.2, 2.0], size=15, h=0.5)
bullets(s, ["Load = invocations/min ÷ 60 × duration (concurrency)", "Timer functions excluded (periodic, not diurnal)", "Chronological split: train 1-10, validation 11-12, test 13-14", "Azure traffic priced with AWS prices"], x=7.3, y=1.9, w=5.6, size=17)
# 6 features and model
fi = FC["feature_importance_q0.9"]
s = slide("Features and model", f"One global LightGBM model per quantile, objective quantile, trained across all heavy functions with load normalised per function. Horizon: next hour, decided at the start of the hour. Features come only from before the decision time: last hour and last six hours mean, max and 90th percentile; same hour yesterday; hour of day; trigger; and the function's level. The safety margin was chosen on validation by minimising decision cost, and came out as {FC['margin_chosen']:.1f}, so none. Baselines: persistence, rolling six-hour quantile, seasonal naive and static training quantile. Accuracy: LightGBM is best at the upper quantiles and tracks coverage; persistence is better at the median. The decision-cost result is on the next slide.", "LightGBM quantile regression, next-hour horizon")
bullets(s, ["Targets: q in {0.5, 0.7, 0.8, 0.9, 0.95, 0.99} of next-hour per-minute concurrency", "Features: last 60 / 360 min (mean, max, q90), same hour yesterday, hour of day, trigger, app level", "Margin multiplier chosen on validation by decision cost", "Baselines: persistence, rolling 6 h, seasonal naive, static"], w=6.4, size=17)
pic(s, "quantile_coverage.png", 7.2, 1.7, w=5.8)
# 7 architecture
s = slide("Architecture", "A local controller does two jobs. Every decision interval it takes the forecast quantile at q-star and sets the ASG desired capacity. Every minute it reads the observed request rate and sets the ALB weights: the EC2 share is the minimum of one and healthy capacity over observed rate, in integer weights from 0 to 999. The ALB splits by percentage, not overflow-only, so the live system leaks relative to the simulator; we measure that. Everything is tagged project equals costscale, default VPC, no NAT gateway, no RDS, and one make down destroys it and verifies with the tagging API.", "Forecast sets instances, observed load sets weights")
pic(s, "architecture.png", 0.4, 1.6, w=8.7)
bullets(s, ["ALB weights 0-999, per minute", "ASG desired capacity, per hour", "Lambda 1769 MB (1 vCPU) as overflow", "make up / make down, tagged, verified"], x=9.3, y=2.0, w=3.8, size=16)
# 8 simulation results
s = slide("Simulation results", f"Test days 13 and 14 scaled to 30 days, EC2 base, container capacity measured on AWS. Hybrid rolling costs {K(T['6_hybrid_rolling']['cost'])}, {100*(1-T['6_hybrid_rolling']['cost']/T['1_all_lambda']['cost']):.0f} percent below all-Lambda and {100*(1-T['6_hybrid_rolling']['cost']/r7['cost']):.0f} percent below reactive at u 0.7, which also drops {r7['unserved_pct']:.1f} percent of load. Plainly: the ML forecaster is only {abs(FC['ml_vs_rolling_pct']):.2f} percent {'cheaper' if FC['ml_vs_rolling_pct']<0 else 'more expensive'} than the rolling six-hour quantile on test, and the whole headroom to the oracle is {FC['headroom_rolling_to_oracle_pct']:.1f} percent. Energy, by an estimate based on the Cloud Carbon Footprint method, is lowest for all-Lambda: the hybrid saves money, not energy. Under the earlier assumed capacity of 16 the saving was larger; measured capacity is {SIM['capacity_units']['ec2']:.1f}.", "The hybrid wins on cost and unserved load; ML adds almost nothing over rolling")
pic(s, "cost_comparison.png", 0.4, 1.6, w=8.4)
bullets(s, [f"Hybrid-rolling {K(T['6_hybrid_rolling']['cost'])} vs all-Lambda {K(T['1_all_lambda']['cost'])}", f"ML vs rolling: {FC['ml_vs_rolling_pct']:+.2f}% (noise level)", f"Headroom to oracle: {FC['headroom_rolling_to_oracle_pct']:.1f}%", f"Energy (estimate): hybrid {T['6_hybrid_rolling']['energy_wh']/1000:,.0f} kWh vs all-Lambda {T['1_all_lambda']['energy_wh']/1000:,.0f} kWh"], x=8.9, y=1.9, w=4.2, size=16)
# 9 live results
if LIVE:
    A, B, C, D = (LIVE.get(k) for k in ("A_reactive", "B_static", "C_ml", "D_ml_spike"))
    rows = [["Arm", "Errors", "Lambda", "p95 ms", "Cost"]]
    for key, nm in (("A_reactive", "A reactive"), ("B_static", "B static"), ("C_ml", "C ML"), ("D_ml_spike", "D ML + spike")):
        v = LIVE.get(key)
        if v: rows.append([nm, f"{100*v['error_rate']:.1f}%", f"{100*v['lambda_share']:.0f}%", f"{v['latency_ms']['all']['p95']:.0f}", f"${v['total_cost_cw']:.2f}"])
    s = slide("Live results on AWS", "Four arms on the same trace window, scaled to 28 requests per second and time-compressed, run from a laptop in Mumbai against the load balancer in ap-south-1. Arm A is ASG target tracking with everything on EC2; B and C use the controller with the static and ML forecasts; D is C with a 4x spike in the fourth hour. The table shows errors, share served by Lambda, p95 latency and cost computed from instance time and CloudWatch Lambda metrics. Each arm was run once, so this supports the design but is not a statistical ranking. Compare to the simulator replay on the report's reconciliation table.", "Four arms, same trace, run once each")
    pic(s, "live_timelines.png", 0.4, 1.5, w=7.6)
    table(s, rows, 8.3, 1.9, 4.7, [1.5, 0.8, 0.8, 0.8, 0.8], size=12, h=0.45)
    bullets(s, [f"Weight propagation about {sum(CAL['weight_propagation_s'])/3:.0f} s", f"Lambda cold start +{CAL['lambda_cold_warm_v2']['cold_overhead_ms_median']:.0f} ms", f"Launch to healthy {CAL['launch_to_healthy_s'][0]:.0f}-{CAL['launch_to_healthy_s'][1]:.0f} s"], x=8.3, y=4.5, w=4.7, size=14)
else:
    s = slide("Live results on AWS", "Live results pending.", "pending")
# 10 prediction wrong
sp = FC["spike_test_ec2"]
s = slide("What if the prediction is wrong?", f"Under-prediction is cheap by design: the excess goes to Lambda, nothing is dropped. Over-prediction wastes f per idle unit-minute until the next decision. A reactive controller is wrong in the damaging way. In simulation we multiplied one hour per function by 3, 4 or 5 after the forecast was made: the hybrids served everything; a x5 spike adds {K(sp['5']['hybrid_rolling']['cost']-T['6_hybrid_rolling']['cost'])} to {K(sp['5']['hybrid_ML']['cost']-T['7_hybrid_ML']['cost'])} per month. The ML forecaster over-reacts after a spike because its features include the spiked hour; a production system would clip forecasts after an anomaly. Live arm D repeated this on AWS.", "Under-prediction goes to Lambda; nothing is dropped")
pic(s, "spike.png", 0.4, 1.8, w=8.5)
bullets(s, ["Too small base: overflow to Lambda, higher cost, higher latency on cold start", "Too large base: idle containers, wasted money only", "Reactive: under-provisions and drops load", f"x5 spike: +{K(sp['5']['hybrid_rolling']['cost']-T['6_hybrid_rolling']['cost'])} (rolling) / +{K(sp['5']['hybrid_ML']['cost']-T['7_hybrid_ML']['cost'])} (ML) per month, 0% dropped"], x=9.1, y=1.9, w=4.0, size=15)
# 11 limitations and future
s = slide("Limitations and future work", "Be upfront. Azure traffic is priced with AWS prices. Capacity is measured for one synthetic request on one instance type from a coarse grid. Live arms are short, time-compressed, single runs from one laptop. Energy is an estimate. AWS predictive scaling was not run live because it needs 24 hours of history. The hybrid advantage depends on measured capacity and Lambda price. Future work: storage tiering, network cost, a Fargate base, multi-region, anomaly-aware forecasting, and overflow-only routing.", "Honest boundaries")
bullets(s, ["Azure traffic, AWS prices", "One request type, one instance type, coarse capacity grid", "Short, compressed, single-run live arms", "Energy is an estimate; predictive scaling not run live", "Advantage depends on capacity and Lambda price"], w=6.0, size=16)
bullets(s, ["Storage tiering and network cost (out of scope here)", "Fargate base, measured and run live", "Multi-region placement", "Anomaly-aware forecasts", "Overflow-only routing instead of weights"], x=7.0, w=5.8, size=16)
# 12 conclusion
s = slide("Conclusion", f"Summarise. The newsvendor rule gives a principled base size and matches brute force. In simulation the hybrid serves all load at {K(T['6_hybrid_rolling']['cost'])} per month against {K(T['1_all_lambda']['cost'])} all-Lambda and {K(r7['cost'])} reactive with {r7['unserved_pct']:.1f} percent dropped. The ML forecaster adds almost nothing over a rolling quantile, and by an estimate the hybrid does not save energy. The live AWS runs are reconciled with a simulator replay. Thank the audience and invite questions.", "A principled base, an honest result")
bullets(s, ["Newsvendor rule sets the base; matches brute force", f"Hybrid: {K(T['6_hybrid_rolling']['cost'])}/month, 0% dropped, vs {K(T['1_all_lambda']['cost'])} all-Lambda and reactive {K(r7['cost'])} with {r7['unserved_pct']:.1f}% dropped", "ML forecaster ≈ rolling quantile (reported plainly)", "Hybrid saves cost, not energy (estimate)", "Live AWS demo reconciled with simulator replay"], w=12, size=20)
# 13 backup: cost model
P = R("prices.json")["values"]
s = slide("Backup: cost-model assumptions", "Prices from the AWS Price List API for ap-south-1, retrieved on the date shown. Azure traffic is priced with AWS prices. Lambda memory is the larger of 128 MB and the application's memory. Container capacity is the measured value for EC2; the Fargate capacity is assumed.", "Backup")
table(s, [["Item", "Value"], ["Lambda duration", f"${P['lambda_gbs']:.7f} per GB-s"], ["Lambda requests", f"${P['lambda_req']:.7f} per request"], ["EC2 t3.medium", f"${P['t3_medium_h']:.4f} per hour"], ["Fargate vCPU / GB", f"${P['fargate_vcpu_h']:.5f} / ${P['fargate_gb_h']:.6f} per hour"], ["EC2 capacity unit", f"{SIM['capacity_units']['ec2']:.1f} concurrent requests (measured; assumed 16)"], ["Fargate capacity unit", "4 concurrent requests (assumed)"], ["Prices retrieved", R("prices.json")["retrieved"]]], 0.6, 1.8, 9.5, [3.5, 6.0], size=15, h=0.5)
# 14 backup: sensitivity
sens = SIM["sensitivity_test"]
s = slide("Backup: sensitivity", "Hybrid-rolling cost relative to all-Lambda under each scenario. The advantage is large when containers carry more load per unit or Lambda is dearer, and small when capacity per unit is halved or Lambda is half price. The EC2 discount scenario is hypothetical, not an AWS-quoted figure.", "Backup")
pic(s, "sensitivity.png", 0.5, 1.6, w=8.4)
bullets(s, [f"Capacity x0.5: saves {100*(1-sens['cap_x0.5']['6_hybrid_rolling']['cost']/sens['cap_x0.5']['1_all_lambda']['cost']):.0f}% vs all-Lambda", f"Lambda price x0.5: saves {100*(1-sens['lambda_price_x0.5']['6_hybrid_rolling']['cost']/sens['lambda_price_x0.5']['1_all_lambda']['cost']):.0f}%", f"Assumed capacity 16: {K(SIM16['ec2']['test']['6_hybrid_rolling']['cost'])} vs {K(SIM16['ec2']['test']['1_all_lambda']['cost'])}"], x=9.2, y=2.0, w=3.9, size=15)
(ROOT / "ppt").mkdir(exist_ok=True); prs.save(ROOT / "ppt/CostScale.pptx"); print("slides", len(prs.slides))
